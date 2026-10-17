"""DOCX processor for eDocAPI."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument

from edocapi.exceptions import ConversionError, InvalidDocument
from edocapi.processors.base import BaseProcessor, register_processor

logger = logging.getLogger("edocapi.processors.docx")


@register_processor("docx")
class DocxProcessor(BaseProcessor):
    supported_input_types = {"docx"}
    supported_output_types = {"pdf", "txt", "html"}

    def __init__(self, path: Path, *, temp_storage: Any = None) -> None:
        super().__init__(path, temp_storage=temp_storage)
        try:
            self._doc = DocxDocument(str(path))
        except Exception as exc:
            raise InvalidDocument(f"Invalid or corrupted DOCX: {exc}") from exc

    def to_text(self) -> str:
        paragraphs = [p.text for p in self._doc.paragraphs if p.text.strip()]
        # Also extract tables
        for table in self._doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                paragraphs.append("\t".join(cells))
        return "\n\n".join(paragraphs)

    def to_html(self) -> str:
        import html as html_mod
        parts = ['<!DOCTYPE html><html><head><meta charset="utf-8"></head><body>']
        for p in self._doc.paragraphs:
            style = p.style.name if p.style else ""
            text = html_mod.escape(p.text)
            if "Heading 1" in style:
                parts.append(f"<h1>{text}</h1>")
            elif "Heading 2" in style:
                parts.append(f"<h2>{text}</h2>")
            elif "Heading 3" in style:
                parts.append(f"<h3>{text}</h3>")
            else:
                parts.append(f"<p>{text}</p>")
        for table in self._doc.tables:
            parts.append("<table border='1'>")
            for row in table.rows:
                parts.append("<tr>")
                for cell in row.cells:
                    parts.append(f"<td>{html_mod.escape(cell.text)}</td>")
                parts.append("</tr>")
            parts.append("</table>")
        parts.append("</body></html>")
        return "\n".join(parts)

    def to_pdf(self) -> Path:
        """Render DOCX into a paginated, styled PDF with visible table borders."""
        try:
            from edocapi.processors.docx_reportlab import (
                convert_docx_with_reportlab,
                new_output_path,
            )
        except ImportError:
            convert_docx_with_reportlab = None
        if convert_docx_with_reportlab is not None:
            reportlab_out = new_output_path(self)
            try:
                result = convert_docx_with_reportlab(self, reportlab_out)
                if result.is_file() and result.stat().st_size:
                    return result
                raise ConversionError("DOCX renderer produced an empty PDF.")
            except Exception as exc:
                self._discard_temp(reportlab_out)
                raise ConversionError("DOCX to PDF conversion failed.") from exc

        html_content = self.to_html()
        try:
            from weasyprint import HTML
        except ImportError:
            from edocapi.processors.simple_pdf import structured_docx_to_pdf
            return structured_docx_to_pdf(self)
        from edocapi.processors.safe_fetch import blocked_url_fetcher
        out = self._temp_path(suffix=".pdf", prefix="docx_")
        try:
            HTML(string=html_content, url_fetcher=blocked_url_fetcher).write_pdf(str(out))
            return out
        except Exception as exc:
            self._discard_temp(out)
            raise ConversionError("DOCX to PDF conversion failed.") from exc

    def info(self) -> dict[str, Any]:
        base = super().info()
        base.update({
            "paragraphs": len(self._doc.paragraphs),
            "tables": len(self._doc.tables),
            "sections": len(self._doc.sections),
        })
        core = self._doc.core_properties
        base["metadata"] = {
            "author": core.author or "",
            "title": core.title or "",
            "created": str(core.created) if core.created else "",
            "modified": str(core.modified) if core.modified else "",
        }
        return base
